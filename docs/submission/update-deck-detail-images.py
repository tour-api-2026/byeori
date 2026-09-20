"""서비스 소개(4쪽)의 상세 이미지를 다섯 장으로 바꾼다.

낡은 루트 탐색 캡처를 추천 코스 14개가 보이는 화면으로 바꾸고, AI 하루 코스 화면을 더한다.
표 안에 이미지를 얹는 양식이라 좌표로 배치한다(가로 폭 안에서 다섯 장을 고르게 나눈다).
"""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Inches

SHOTS = Path(__file__).parent / 'detail'
NEW = ['x-explore.png', 'f4-3-ai-preview.png']   # 4번째 교체 + 5번째 추가
W, H = Inches(1.8), Inches(3.2)
LEFT, RIGHT = Inches(2.35), Inches(12.75)         # 배치 가능한 가로 구간

prs = Presentation(sys.argv[1])
slide = prs.slides[3]
pics = [sh for sh in slide.shapes if sh.shape_type == 13]
# 대표 이미지(썸네일)는 윗줄, 상세 이미지는 아랫줄. 가로 위치로 나누면 썸네일이 섞인다.
row2_top = min(sh.top for sh in pics) + Inches(1.0)
detail = sorted([sh for sh in pics if sh.top > row2_top], key=lambda s: s.left)
if len(detail) != 4:
    sys.exit(f'상세 이미지가 4장이 아님: {len(detail)}')

top = detail[0].top
for sh in detail[-1:]:                            # 마지막(낡은 루트 탐색) 제거
    sh._element.getparent().remove(sh._element)
keep = detail[:-1]

# 이야기 순서대로 둔다: 홈 → 전통 행사 → 장소 상세 → 추천 코스 → AI 코스
ordered = list(keep) + [slide.shapes.add_picture(str(SHOTS / name), LEFT, top, width=W, height=H)
                        for name in NEW]
gap = (RIGHT - LEFT - W * len(ordered)) // (len(ordered) - 1)
for i, sh in enumerate(ordered):
    sh.left, sh.top, sh.width, sh.height = LEFT + i * (W + gap), top, W, H
pics = ordered

prs.save(sys.argv[2])
print('저장:', sys.argv[2], '· 상세 이미지', len(pics), '장')
