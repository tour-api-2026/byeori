"""흐름도 5쪽에 실제 서비스 캡처를 넣는다(양식은 4칸 표에 캡처+단계 설명을 요구한다)."""
import copy, sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.text.text import _Paragraph

SHOTS = Path('/tmp/claude-1000/-home-hidi-dev-byeori/4cd49158-baf2-4727-b24f-9c8e916c8cf3/scratchpad/crops')
EMU, BODY = 914400, RGBColor(0x1A, 0x1A, 0x1F)

FLOWS = [
    [('f1-1-home',      '홈 화면에 전통 테마 행사가 바로 보인다'),
     ('f1-2-section',   '진행 중인 행사를 포스터와 함께 노출한다'),
     ('f1-3-list',      '전체보기에서 전통 행사만 모아 본다'),
     ('f1-4-detail',    '행사 상세에서 기간·장소를 보고 원문으로 이어진다')],
    [('f2-1-map',       '현재 위치를 중심으로 주변 장소를 표시한다'),
     ('f2-2-filter',    '카테고리로 보고 싶은 유형만 남긴다'),
     ('f2-3-card',      '마커를 누르면 요약 카드가 뜬다'),
     ('f2-4-cluster',   '축소하면 가까운 장소끼리 묶여 개수로 보인다')],
    [('f3-1-search',    '문화·체험·맛집·카페·전통시장·공예·한옥으로 나누어 찾는다'),
     ('f3-2-category',  '전통시장처럼 전통문화에 맞춘 분류로 좁힌다'),
     ('f3-3-detail',    '한복 착용 혜택과 소개글을 함께 보여준다'),
     ('f3-4-live',      '이용시간·휴무일·문의처를 실시간 조회하고 출처를 밝힌다')],
    [('f4-1-routes',    '루트 탭에서 여행 일지를 시작한다'),
     ('f4-2-list',      '만든 일지를 한눈에 관리한다'),
     ('f4-3-detail',    '날짜별로 방문지와 시각을 구성한다'),
     ('f4-4-route',     '이동 경로와 총 거리·소요 시간을 지도에서 확인한다')],
    [('f5-1-write',     '별점과 후기를 남긴다'),
     ('f5-2-rating',    '작성한 리뷰가 장소 평점에 즉시 반영된다'),
     ('f5-3-myreviews', '내가 쓴 리뷰를 모아 수정·삭제한다'),
     ('f5-4-blocked',   '차단한 이용자를 관리한다 (신고는 각 화면에서)')],
]


def set_cell(cell, text, size):
    tf = cell.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    p0.runs[0].text = text
    p0.runs[0].font.color.rgb = BODY
    p0.runs[0].font.bold = False
    p0.runs[0].font.size = Pt(size)
    for extra in p0.runs[1:]:
        extra._r.getparent().remove(extra._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def body_table(slide):
    for sh in slide.shapes:
        if sh.has_table and len(sh.table.rows) == 3 and len(sh.table.columns) == 4:
            return sh
    sys.exit('흐름도 표를 찾지 못함')


prs = Presentation(sys.argv[1])
for i, steps in enumerate(FLOWS):
    slide = prs.slides[4 + i]
    for sh in list(slide.shapes):
        if sh.__class__.__name__ == 'Picture':
            sh._element.getparent().remove(sh._element)
    shape = body_table(slide)
    t = shape.table
    for r in (1, 2):
        if t.cell(r, 0).is_merge_origin:
            t.cell(r, 0).split()
    col_w = shape.width / 4
    row1_top = shape.top + t.rows[0].height
    row1_h = t.rows[1].height
    pad = int(0.06 * EMU)
    for c, (name, desc) in enumerate(steps):
        img = SHOTS / f'{name}.png'
        if not img.exists():
            sys.exit(f'캡처 없음: {img}')
        iw, ih = Image.open(img).size
        h = row1_h - pad * 2
        w = int(h * iw / ih)
        left = int(shape.left + c * col_w + (col_w - w) / 2)
        slide.shapes.add_picture(str(img), left, int(row1_top + pad), width=w, height=h)
        set_cell(t.cell(2, c), f'{c + 1}. {desc}', 11)

prs.save(sys.argv[2])
print('저장:', sys.argv[2])
