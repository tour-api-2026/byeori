"""기능설명서에 AI 하루 코스와 추천 코스 확대를 반영한다.

- 3쪽 핵심기능: 여행 일정 → AI 하루 코스
- 6쪽 지도 설명: 목록이 저장 기반으로 바뀐 것을 반영(예전 '실시간 조회' 문구가 남아 있었다)
- 8쪽 핵심기능4: AI 코스 흐름과 캡처 4장
- 11쪽 기타 API: OpenAI 추가
- 12쪽 차별성·발전계획: 한복 항목 → AI 코스, 추천 코스 14개 반영
"""
import copy, sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.text.text import _Paragraph

BODY = RGBColor(0x1A, 0x1A, 0x1F)
EMU = 914400
SHOTS = Path('/tmp/claude-1000/-home-hidi-dev-byeori/4cd49158-baf2-4727-b24f-9c8e916c8cf3/scratchpad/crops')


def style(run, size):
    run.font.color.rgb = BODY
    run.font.bold = False
    if size:
        run.font.size = Pt(size)


def set_cell(cell, text, size=None):
    """셀 서식(첫 문단)을 유지한 채 여러 줄 텍스트로 바꾼다."""
    tf = cell.text_frame
    lines = text.split('\n')
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    size = size or (p0.runs[0].font.size.pt if p0.runs[0].font.size else None)
    p0.runs[0].text = lines[0]
    style(p0.runs[0], size)
    for extra in p0.runs[1:]:
        extra._r.getparent().remove(extra._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for line in lines[1:]:
        newp = copy.deepcopy(p0._p)
        p0._p.getparent().append(newp)
        para = _Paragraph(newp, tf)
        para.runs[0].text = line
        style(para.runs[0], size)


def tables(slide):
    return [sh for sh in slide.shapes if sh.has_table]


def replace_block(text, start, end, new):
    """start 로 시작하는 항목부터 end 직전까지를 new 로 바꾼다(end=None 이면 끝까지)."""
    i = text.index(start)
    j = text.index(end, i) if end else len(text)
    return text[:i] + new + text[j:]


prs = Presentation(sys.argv[1])

# 3쪽 핵심기능
cell = tables(prs.slides[2])[0].table.cell(0, 1)
t = cell.text
t = replace_block(t, '- 여행 일정 생성 및 경로 안내 기능', '- 리뷰 및 이용자 보호 기능',
    '- AI 하루 코스 및 경로 안내 기능\n'
    '   지역·테마만 고르면 AI가 벼리의 장소·행사로 하루 코스를 짜고, 저장하면\n'
    '   실제 이동 경로와 소요 시간 확인. 전국 12개 권역 추천 코스도 제공\n')
set_cell(cell, t.rstrip('\n'))

# 6쪽 지도 설명
for sh in tables(prs.slides[5]):
    tb = sh.table
    if tb.cell(0, 0).text.strip() == '핵심 기능2':
        set_cell(tb.cell(1, 1),
                 '저장된 장소 가운데 지도가 보고 있는 영역만 불러와 표시하고, 축소하면 가까운 장소끼리 '
                 '묶어 개수로 보여줍니다. 위치 정보는 지도 이동에만 쓰고 서버로 전송하지 않습니다.')

# 8쪽 핵심기능4 — 제목·설명·흐름도 캡처
slide = prs.slides[7]
STEPS = [('f4-1-ai-card', '루트 탭의 AI 카드로 시작한다'),
         ('f4-2-ai-form', '지역·테마·날짜와 원하는 조건을 고른다'),
         ('f4-3-ai-preview', 'AI가 벼리의 장소 안에서 코스를 짜고 요청대로 다듬는다'),
         ('f4-4-ai-route', '저장한 코스의 이동 경로와 소요 시간을 지도에서 본다')]
for sh in tables(slide):
    tb = sh.table
    if tb.cell(0, 0).text.strip() == '핵심 기능4':
        set_cell(tb.cell(0, 1), 'AI 하루 코스와 경로 안내')
        set_cell(tb.cell(1, 1),
                 '지역·테마·날짜와 "아이와 함께" 같은 요청을 받아, 서버가 벼리의 장소·행사에서 후보와 하루 틀을 '
                 '정하고 AI가 그 안에서 코스를 짭니다. 결과는 "점심을 바꿔 주세요"처럼 말해 다듬을 수 있고, '
                 '저장하면 일정이 되어 이동 경로와 총 거리·소요 시간을 지도에서 확인합니다.')
for pic in [s for s in slide.shapes if s.__class__.__name__ == 'Picture']:
    pic._element.getparent().remove(pic._element)
shape = next(sh for sh in tables(slide) if len(sh.table.rows) == 3 and len(sh.table.columns) == 4)
tb = shape.table
col_w = shape.width / 4
row1_top = shape.top + tb.rows[0].height
row1_h = tb.rows[1].height
pad = int(0.06 * EMU)
for c, (name, desc) in enumerate(STEPS):
    img = SHOTS / f'{name}.png'
    iw, ih = Image.open(img).size
    h = row1_h - pad * 2
    w = int(h * iw / ih)
    left = int(shape.left + c * col_w + (col_w - w) / 2)
    slide.shapes.add_picture(str(img), left, int(row1_top + pad), width=w, height=h)
    set_cell(tb.cell(2, c), f'{c + 1}. {desc}', 11)

# 11쪽 기타 API
tb = tables(prs.slides[10])[0].table
set_cell(tb.cell(4, 2), '카카오맵 / 카카오모빌리티 API, OpenAI API')
set_cell(tb.cell(5, 2), '지도 표시와 장소 키워드 검색, 여행 일정의 이동 경로·소요 시간 계산에 카카오를 활용. '
                        'AI 하루 코스 생성에 OpenAI를 활용(벼리가 뽑은 후보 안에서만 선택)')

# 12쪽 차별성·발전계획
tb = tables(prs.slides[11])[0].table
t = tb.cell(0, 1).text
t = replace_block(t, '6. 한복 착용 혜택 정보', None,
    '6. 지어내지 않는 AI 하루 코스\n'
    '   AI에게 장소를 만들게 하지 않습니다. 서버가 벼리의 장소·행사에서 후보와 하루 틀을 정하고,\n'
    '   AI는 그 안에서 고르고 이유만 씁니다. 응답은 후보와 대조해 없는 장소가 화면에 나오지 않습니다.')
set_cell(tb.cell(0, 1), t)
t = tb.cell(1, 1).text
t = replace_block(t, '2. 지역별 심화 큐레이션', '3. 무장애 여행 정보 결합',
    '2. 지역별 심화 큐레이션\n'
    '   전국 12개 권역 추천 코스를 시작으로, 전주 한옥마을·안동 하회마을 등 지역 단위 전통문화\n'
    '   테마 코스를 늘려 지역 관광 활성화에 기여합니다.\n')
t = replace_block(t, '4. 개인화 추천', '5. 모바일 앱 배포',
    '4. 개인화 추천\n'
    '   AI 하루 코스에 즐겨찾기와 방문 이력을 반영해 관심 지역·유형에 맞는 코스를 만듭니다.\n')
set_cell(tb.cell(1, 1), t)

prs.save(sys.argv[2])
print('저장:', sys.argv[2])
