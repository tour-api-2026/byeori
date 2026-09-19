"""기능설명서를 현재 구현에 맞춘다.

목록(지도·검색·지역 탐색)이 저장 기반으로, 상세만 실시간으로 바뀌었다.
수집 범위도 6개 유형으로 넓어졌다. 이미지가 든 산출물이라 제자리에서 고친다.
"""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.text.text import _Paragraph

BODY = RGBColor(0x1A, 0x1A, 0x1F)


def style(run, size):
    run.font.color.rgb = BODY
    run.font.bold = False
    if size:
        run.font.size = Pt(size)


def set_cell(cell, text, size=None):
    tf = cell.text_frame
    lines = text.split('\n')
    p0 = tf.paragraphs[0]
    if not p0.runs:
        p0.add_run()
    p0.runs[0].text = lines[0]
    style(p0.runs[0], size)
    for extra in p0.runs[1:]:
        extra._r.getparent().remove(extra._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for line in lines[1:]:
        newp = copy.deepcopy(p0._p)
        p0._p.getparent().append(newp)
        para = _Paragraph(newp, tf)
        para.runs[0].text = line
        style(para.runs[0], size)


def table_of(slide, rows, cols):
    for sh in slide.shapes:
        if sh.has_table and len(sh.table.rows) == rows and len(sh.table.columns) == cols:
            return sh.table
    sys.exit(f'표를 찾지 못함 {rows}x{cols}')


prs = Presentation(sys.argv[1])

# ── 3쪽 핵심기능 ────────────────────────────────────────
set_cell(table_of(prs.slides[2], 1, 2).cell(0, 1),
         '- 전통 테마 행사 큐레이션 기능\n'
         '   공연·축제 데이터에서 전통을 주제로 한 행사만 자동 분류해 제공하고, 행사마다\n'
         '   포스터·기간·장소와 원문 페이지를 연결\n'
         '- 위치 기반 지도 탐색 기능\n'
         '   현재 위치 중심으로 주변 장소를 표시하고, 축소하면 가까운 것끼리 묶어 개수로 보여줌\n'
         '- 명칭 검색 및 지역별 탐색 기능\n'
         '   문화·체험·맛집·카페·전통시장·공예·한옥스테이 일곱 갈래로 나누어 탐색\n'
         '- 장소 상세 정보 제공 기능\n'
         '   공사 OpenAPI를 실시간 조회해 소개글·이용시간·휴무일·문의처를 표시하고,\n'
         '   진행 중인 행사를 함께 제공\n'
         '- 여행 일정 생성 및 경로 안내 기능\n'
         '   방문지를 담아 하루 일정을 구성하고 실제 이동 경로와 소요 시간 확인\n'
         '- 리뷰 및 이용자 보호 기능\n'
         '   방문 후기 작성, 부적절한 콘텐츠 신고, 특정 이용자 차단')

# ── 10쪽 공사 OpenAPI 목록 (실제로 부르는 것만) ──────────
APIS = [
    ('한국관광공사 국문 관광정보 서비스 (detailCommon2 / detailIntro2)',
     '장소·행사 상세를 열 때마다 실시간 호출하여 소개글·이용시간·휴무일·문의처·주차·홈페이지를 조회. '
     '저장 목록에 없는 장소도 콘텐츠 ID로 동일하게 조회'),
    ('한국관광공사 국문 관광정보 서비스 (areaBasedSyncList2)',
     '공사가 로컬 저장용으로 제공하는 동기화 API. 1일 1회 직전 수집 이후 변경분만 수신하며, '
     'showflag로 공사가 내린 콘텐츠를 함께 비노출 처리'),
    ('한국관광공사 국문 관광정보 서비스 (searchFestival2)',
     '행사·축제 정보를 1일 1회 수집. 전통 테마 행사 자동 분류의 기반으로 활용'),
    ('한국관광공사 국문 관광정보 서비스 (areaBasedList2)',
     '초기 장소 적재에 활용. 관광지·문화시설·음식점·레포츠·쇼핑·숙박 여섯 유형을 수집한 뒤 '
     '신분류체계로 전통문화와 닿는 것만 선별'),
    ('한국관광공사 국문 관광정보 서비스 (ldongCode2)',
     '법정동·시군구 코드 조회. 전국 단위 수집 시 지역을 나누는 기준으로 활용'),
]
t = table_of(prs.slides[9], 10, 3)
for i, (name, desc) in enumerate(APIS):
    set_cell(t.cell(i * 2, 2), name, size=12)
    set_cell(t.cell(i * 2 + 1, 2), desc, size=11)

# ── 12쪽 차별성 ─────────────────────────────────────────
set_cell(table_of(prs.slides[11], 2, 2).cell(0, 1),
         '1. 장소 정보와 공연 정보의 결합\n'
         '   대부분의 관광 서비스는 장소 또는 공연 한쪽만 다룹니다. 벼리는 한국관광공사 OpenAPI의 장소\n'
         '   데이터와 공연 데이터를 연결해, 장소 상세에서 "여기서 지금 무엇이 열리는지"를 함께 보여줍니다.\n'
         '2. 전통 행사 자동 분류\n'
         '   공공데이터는 전통 행사를 따로 구분해 주지 않습니다. 장르 코드만으로는 놓치는 행사가 많아\n'
         '   제목·주최기관을 함께 검사하는 규칙으로 선별했습니다.\n'
         '3. 전통문화에 맞춘 분류 체계\n'
         '   공사 유형 코드로는 카페와 맛집이 한 덩어리이고, 쇼핑·숙박은 대부분 전통문화와 무관합니다.\n'
         '   신분류체계를 파고들어 전통시장·공예·한옥스테이를 따로 골라냈습니다.\n'
         '4. 목록은 빠르게, 상세는 최신으로\n'
         '   명칭·주소·좌표는 잘 변하지 않아 저장분으로 즉시 보여주고, 이용시간·휴무일처럼 자주 바뀌는\n'
         '   정보는 상세를 열 때 공사 OpenAPI를 실시간 조회합니다.\n'
         '5. 공사 데이터와 어긋나지 않는 운영\n'
         '   공사가 제공하는 동기화 API로 변경분만 받고, 공사가 내린 콘텐츠는 showflag로 확인해\n'
         '   즉시 비노출 처리합니다. 사라진 장소가 서비스에 남지 않습니다.\n'
         '6. 한복 착용 혜택 정보\n'
         '   한복 착용 시 혜택이 있는 장소를 별도로 표시해 전통문화 체험이 하나의 동선으로 이어지게 했습니다.', size=10.5)

prs.save(sys.argv[2])
print('저장:', sys.argv[2])
